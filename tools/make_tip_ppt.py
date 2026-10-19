# -*- coding: utf-8 -*-
"""하루 한 팁 Day 08·09 연습용 PPT — download/하루팁_PPT연습.pptx (가상 수업 · 20장)
   B/W로 정답 가리기, Shift+F5, 숫자+Enter(12번으로), Ctrl+P 판서를 직접 해 본다.
   사용: python tools/make_tip_ppt.py"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "download", "하루팁_PPT연습.pptx")
INK, GREEN, PINK, MUTE = RGBColor(0x11, 0x11, 0x11), RGBColor(0x16, 0xA3, 0x4A), RGBColor(0xED, 0x1A, 0xA0), RGBColor(0x70, 0x70, 0x72)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size, color=INK, bold=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
    return tf


def num(slide, n):
    text(slide, 12.0, 6.85, 1.1, 0.4, f"{n} / 20", 14, MUTE)


def box(slide, x, y, w, h, fill, label, size=24, color=RGBColor(255, 255, 255)):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    tf = sh.text_frame
    tf.text = label
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
            r.font.color.rgb = color
    return sh


TOPICS = {2: "씨앗 관찰하기", 3: "씨앗의 생김새 비교", 4: "씨앗 심기", 5: "물 주기 계획", 6: "싹이 트는 모습", 7: "떡잎과 본잎",
          8: "줄기가 자라요", 9: "관찰 기록 쓰는 법", 10: "모둠 관찰 일지", 11: "중간 정리", 13: "실험: 물이 없으면?",
          14: "실험: 추우면?", 15: "실험 결과 정리", 18: "꽃과 열매", 19: "한살이 순서 정리"}

for n in range(1, 21):
    s = prs.slides.add_slide(BLANK)
    if n == 1:
        text(s, 0.9, 1.6, 11, 0.6, "과학 4학년 (가상 수업) · 하루 한 팁 PPT 연습", 22, GREEN, True)
        text(s, 0.9, 2.3, 11, 1.4, "식물의 한살이", 66, INK, True)
        text(s, 0.9, 4.0, 11.5, 2.6,
             "따라 해 보기\n"
             "· F5 = 처음부터 / Shift + F5 = 지금 슬라이드부터\n"
             "· 슬라이드쇼 중 12 Enter → 12번 슬라이드로 바로\n"
             "· 16번 퀴즈에서 B = 검은 화면, W = 흰 화면 (아무 키나 누르면 돌아와요)\n"
             "· Ctrl + P = 펜으로 판서, E = 지우기, Ctrl + A = 화살표", 20, MUTE)
    elif n == 12:
        text(s, 0.9, 1.2, 8, 0.6, "과학 4학년 (가상 수업)", 22, GREEN, True)
        text(s, 0.9, 1.9, 8, 1.2, "12. 싹이 트는 조건", 54, INK, True)
        text(s, 0.9, 3.3, 8, 0.8, "물 · 온도 · 공기 중 무엇이 필요할까?", 28, MUTE)
        for i, (lab, c) in enumerate((("물", RGBColor(0x25, 0x63, 0xEB)), ("온도", RGBColor(0xF5, 0x9E, 0x0B)), ("공기", RGBColor(0x0E, 0xA5, 0xE9)))):
            box(s, 0.9 + i * 2.6, 4.6, 2.3, 1.1, c, lab, 28)
        text(s, 9.3, 2.2, 3.6, 3.0, "여기서 Ctrl + P로\n동그라미를 그려 보세요", 20, PINK, True)
    elif n == 16:
        text(s, 0.9, 1.0, 11, 0.6, "퀴즈 (가상)", 24, RGBColor(0x25, 0x63, 0xEB), True)
        text(s, 0.9, 1.7, 11.5, 1.2, "씨가 싹 트는 데 꼭 필요하지 않은 것은?", 44, INK, True)
        text(s, 0.9, 3.2, 11.5, 1.0, "① 물　　② 알맞은 온도　　③ 햇빛　　④ 공기", 32, INK)
        text(s, 0.9, 5.2, 11.5, 0.8, "다음 장에 정답이 있어요. 넘기기 전에 B를 눌러 화면을 가려 보세요.", 20, PINK, True)
    elif n == 17:
        text(s, 0.9, 1.7, 11.5, 1.2, "정답: ③ 햇빛", 60, RGBColor(0xB9, 0x1C, 0x1C), True)
        text(s, 0.9, 3.3, 11.5, 2.0, "씨는 흙 속 어두운 곳에서도 싹이 터요.\n싹이 튼 뒤 자라는 데에는 햇빛이 필요해요.", 28, INK)
    elif n == 20:
        text(s, 0.9, 2.4, 11.5, 1.2, "오늘 배운 단축키", 54, INK, True)
        text(s, 0.9, 3.8, 11.5, 2.0, "Shift + F5 · 숫자 + Enter · B / W · Ctrl + P", 30, PINK, True)
    else:
        title = TOPICS.get(n, f"{n}. 활동")
        text(s, 0.9, 1.2, 11, 0.6, "과학 4학년 (가상 수업)", 20, GREEN, True)
        text(s, 0.9, 1.8, 11.5, 1.2, f"{n}. {title}", 44, INK, True)
        text(s, 0.9, 3.3, 11.5, 2.5, "연습용 빈 슬라이드예요.\n슬라이드쇼 중에 번호를 누르고 Enter로 이동해 보세요.", 24, MUTE)
    num(s, n)

prs.save(OUT)
print("저장:", OUT, f"{os.path.getsize(OUT) // 1024}KB")
